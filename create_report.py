from docx import Document
from docx.shared import Pt

document = Document()

# Title
title = document.add_heading(
    "College Placement Intelligence & Prediction System",
    0
)

document.add_paragraph(
    "Business Intelligence and Predictive Analytics Project"
)

document.add_paragraph(
    "Author: Mahitha"
)

# 1. Introduction
document.add_heading("1. Introduction", level=1)

document.add_paragraph(
    "The College Placement Intelligence & Prediction System is a "
    "Business Intelligence and Predictive Analytics project designed "
    "to analyze historical college placement data. The project "
    "combines data cleaning, exploratory analysis, KPI development, "
    "predictive modeling, and dashboard visualization."
)

# 2. Problem Statement
document.add_heading("2. Problem Statement", level=1)

document.add_paragraph(
    "Educational institutions collect large amounts of student "
    "academic and placement information. However, raw data alone "
    "does not provide decision-ready insights. This project aims "
    "to transform placement data into meaningful business "
    "intelligence that can help understand placement patterns, "
    "salary outcomes, and estimated placement probability."
)

# 3. Objectives
document.add_heading("3. Objectives", level=1)

objectives = [
    "Analyze historical student placement data.",
    "Calculate important placement KPIs.",
    "Explore relationships between student characteristics and placement outcomes.",
    "Analyze salary patterns among placed students.",
    "Build a Logistic Regression model for placement prediction.",
    "Estimate placement probability for individual student profiles.",
    "Present insights through an interactive Streamlit dashboard."
]

for item in objectives:
    document.add_paragraph(item, style="List Bullet")

# 4. Dataset
document.add_heading("4. Dataset Description", level=1)

document.add_paragraph(
    "The project uses the Placement_Data_Full_Class.csv dataset "
    "containing information about 215 students and 15 original "
    "columns."
)

document.add_paragraph(
    "The dataset contains academic performance, educational "
    "background, work experience, MBA specialization, placement "
    "status, and salary information."
)

document.add_paragraph(
    "Important variables include SSC percentage, HSC percentage, "
    "degree percentage, entrance test percentage, MBA percentage, "
    "work experience, specialization, placement status, and salary."
)

# 5. Technologies
document.add_heading("5. Technologies Used", level=1)

technologies = [
    "Python",
    "Pandas",
    "NumPy",
    "Matplotlib",
    "Seaborn",
    "Scikit-learn",
    "Jupyter Notebook",
    "Streamlit",
    "GitHub Codespaces"
]

for item in technologies:
    document.add_paragraph(item, style="List Bullet")

# 6. Data Preprocessing
document.add_heading("6. Data Preprocessing", level=1)

document.add_paragraph(
    "The dataset was loaded using Pandas and inspected for data "
    "types, missing values, and duplicate records. The serial "
    "number column was removed because it does not provide useful "
    "predictive information."
)

document.add_paragraph(
    "Missing salary values were retained because salary is naturally "
    "unavailable for students who were not placed."
)

# 7. KPI Analysis
document.add_heading("7. Business Intelligence and KPI Analysis", level=1)

document.add_paragraph(
    "The project calculates key placement indicators including "
    "total students, placed students, placement rate, average "
    "salary, highest salary, and average academic percentages."
)

document.add_paragraph(
    "These KPIs provide an executive-level overview of the "
    "placement dataset and form the foundation of the dashboard."
)

# 8. Exploratory Analysis
document.add_heading("8. Exploratory and Diagnostic Analysis", level=1)

document.add_paragraph(
    "Placement outcomes were analyzed across work experience, "
    "degree percentage groups, entrance test score groups, and "
    "MBA specialization."
)

document.add_paragraph(
    "The project also analyzes average salary by MBA specialization "
    "and uses a correlation matrix to examine relationships between "
    "numerical variables."
)

document.add_paragraph(
    "These analyses identify associations and patterns in the "
    "historical dataset. They should not be interpreted as proof "
    "of causal relationships."
)

# 9. Predictive Modeling
document.add_heading("9. Predictive Modeling", level=1)

document.add_paragraph(
    "A Logistic Regression model was developed to estimate the "
    "probability that a student would be placed."
)

document.add_paragraph(
    "The model uses pre-placement attributes including academic "
    "percentages, work experience, educational background, and "
    "MBA specialization."
)

document.add_paragraph(
    "Placement status was used as the target variable. Salary was "
    "not used as a predictive feature because it represents an "
    "outcome associated with placement."
)

# 10. Model Evaluation
document.add_heading("10. Model Evaluation", level=1)

document.add_paragraph(
    "The dataset was divided into training and testing subsets. "
    "The model was evaluated using accuracy, precision, recall, "
    "confusion matrix, and classification report."
)

document.add_paragraph(
    "These metrics provide different perspectives on model "
    "performance and help assess how well the model distinguishes "
    "between placed and not-placed students."
)

# 11. Risk Categories
document.add_heading("11. Placement Risk Categorization", level=1)

document.add_paragraph(
    "Predicted placement probabilities are grouped into project-"
    "defined categories. Probabilities of 70% or above are labeled "
    "Lower Risk, probabilities from 40% to below 70% are labeled "
    "Medium Risk, and probabilities below 40% are labeled Higher Risk."
)

document.add_paragraph(
    "These thresholds are defined specifically for this project "
    "and should not be interpreted as universal standards."
)

# 12. Dashboard
document.add_heading("12. Streamlit Dashboard", level=1)

document.add_paragraph(
    "An interactive Streamlit dashboard was developed to present "
    "the project's results in a business-friendly format."
)

dashboard_sections = [
    "Executive Overview",
    "Diagnostic Analysis",
    "Placement Prediction",
    "Salary Analysis",
    "Key Business Insights"
]

for item in dashboard_sections:
    document.add_paragraph(item, style="List Bullet")

# 13. Business Insights
document.add_heading("13. Business Insights", level=1)

document.add_paragraph(
    "The analysis provides evidence about how placement outcomes "
    "vary across different student characteristics. Work experience, "
    "academic performance, entrance test performance, and MBA "
    "specialization can be compared using placement rates."
)

document.add_paragraph(
    "The dashboard converts these analytical results into "
    "decision-support information for understanding placement "
    "patterns."
)

# 14. Recommendations
document.add_heading("14. Recommendations", level=1)

recommendations = [
    "Placement teams can monitor academic and placement KPIs regularly.",
    "Students can be encouraged to develop relevant practical experience.",
    "Academic and entrance-test performance can be monitored to identify areas for additional support.",
    "Placement teams can use historical patterns as one input when planning student support initiatives.",
    "Predictive probabilities should be used as decision-support information rather than guarantees."
]

for item in recommendations:
    document.add_paragraph(item, style="List Bullet")

# 15. Limitations
document.add_heading("15. Limitations", level=1)

limitations = [
    "The dataset contains 215 students and may not represent all colleges or student populations.",
    "The analysis is based on historical data.",
    "Observed associations do not establish causation.",
    "The predictive model should not be treated as a guarantee of placement.",
    "Risk thresholds used in the dashboard are project-defined."
]

for item in limitations:
    document.add_paragraph(item, style="List Bullet")

# 16. Future Scope
document.add_heading("16. Future Scope", level=1)

future_scope = [
    "Use larger and more recent placement datasets.",
    "Compare multiple machine learning algorithms.",
    "Add interactive dashboard filters.",
    "Add additional student and company-level information.",
    "Deploy the dashboard online.",
    "Monitor model performance as new placement data becomes available."
]

for item in future_scope:
    document.add_paragraph(item, style="List Bullet")

# 17. Conclusion
document.add_heading("17. Conclusion", level=1)

document.add_paragraph(
    "The College Placement Intelligence & Prediction System "
    "demonstrates how Business Intelligence and Predictive Analytics "
    "can be combined to transform placement data into useful "
    "decision-support information. The project covers the complete "
    "workflow from data preparation and KPI analysis to predictive "
    "modeling and interactive dashboard development."
)

# Basic font formatting
for paragraph in document.paragraphs:
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(11)

# Save
output_file = "Mahitha_CollegePlacementReport.docx"
document.save(output_file)

print(f"Report created successfully: {output_file}")